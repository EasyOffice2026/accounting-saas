"""Accounting (general ledger). Only logins allowed in acc_settings (finance.manager during testing) can call these."""
import io
import json
from datetime import date
from typing import Optional

from fastapi import APIRouter, Depends, HTTPException, Form
from fastapi.responses import StreamingResponse
from sqlalchemy import func
from sqlalchemy.orm import Session

from app.database import get_db
from app.models.accounting import AccAccount, AccMapping, AccJournal, AccJournalLine
from app.models.branch import Branch
from app.models.channel import PaymentChannel
from app.models.expense import Expense, ExpenseCategory
from app.models.hr import Brand, Employee
from app.models.purchase import PurchaseCategory, Supplier
from app.models.sale import Sale
from app.models.user import User
from app.utils.ledger import (EXPENSE_CATEGORY_DEFAULTS, PURCHASE_CATEGORY_DEFAULTS, SETTING_DEFAULTS, Ledger, ensure_chart,
                              post_journal, reverse_journal, sync_ledger, _r, EPS)
from app.utils.ledger_access import get_accounting_user, get_settings

router = APIRouter(prefix="/api/accounting", tags=["accounting"])

TYPES = ("ASSET", "LIABILITY", "EQUITY", "REVENUE", "EXPENSE")
DEBIT_NATURE = ("ASSET", "EXPENSE")


def _d(v: Optional[str]) -> Optional[date]:
    return date.fromisoformat(v) if v else None


def _brand(db: Session, brand_id: int) -> Brand:
    b = db.query(Brand).filter(Brand.id == brand_id).first()
    if not b:
        raise HTTPException(404, "Company not found")
    return b


def _acc_json(a: AccAccount) -> dict:
    return {"id": a.id, "code": a.code, "name": a.name, "name_ar": a.name_ar, "type": a.type, "subtype": a.subtype,
            "parent_code": a.parent_code, "is_header": bool(a.is_header), "concept": a.concept, "is_active": bool(a.is_active),
            "system": bool(a.system_key)}


def _sums(db: Session, brand_id: int, date_from: Optional[date] = None, date_to: Optional[date] = None,
          branch_id: Optional[int] = None, before: bool = False) -> dict[int, tuple[float, float]]:
    q = db.query(AccJournalLine.account_id, func.coalesce(func.sum(AccJournalLine.debit), 0),
                 func.coalesce(func.sum(AccJournalLine.credit), 0)).join(AccJournal, AccJournal.id == AccJournalLine.journal_id)\
        .filter(AccJournal.brand_id == brand_id)
    if before:
        q = q.filter(AccJournal.date < date_from)
    else:
        if date_from:
            q = q.filter(AccJournal.date >= date_from)
        if date_to:
            q = q.filter(AccJournal.date <= date_to)
    if branch_id:
        q = q.filter(AccJournalLine.branch_id == branch_id)
    return {a: (float(d), float(c)) for a, d, c in q.group_by(AccJournalLine.account_id).all()}


def _accounts(db: Session, brand_id: int) -> list[AccAccount]:
    return db.query(AccAccount).filter(AccAccount.brand_id == brand_id).order_by(AccAccount.code).all()


# ---------- settings / sync ----------

@router.get("/overview")
def overview(db: Session = Depends(get_db), user: User = Depends(get_accounting_user)):
    s = get_settings(db)
    ensure_chart(db)
    brands = db.query(Brand).order_by(Brand.id).all()
    counts = dict(db.query(AccJournal.brand_id, func.count(AccJournal.id)).group_by(AccJournal.brand_id).all())
    users = db.query(User).filter(User.is_active.is_(True)).order_by(User.username).all()
    return {
        "start_date": str(s.start_date) if s.start_date else None,
        "auto_sync": bool(s.auto_sync),
        "user_ids": json.loads(s.user_ids) if s.user_ids else [],
        "users": [{"id": u.id, "username": u.username, "full_name": u.full_name} for u in users if u.role in ("owner", "manager", "accountant")],
        "last_sync_at": s.last_sync_at.isoformat() + "Z" if s.last_sync_at else None,
        "last_sync_result": json.loads(s.last_sync_result) if s.last_sync_result else None,
        "brands": [{"id": b.id, "name_en": b.name_en, "name_ar": b.name_ar, "journals": counts.get(b.id, 0)} for b in brands],
        "branches": [{"id": b.id, "name": b.name, "name_ar": b.name_ar, "brand_id": b.brand_id} for b in db.query(Branch).order_by(Branch.id)],
    }


@router.post("/settings")
def save_settings(start_date: str = Form(""), auto_sync: bool = Form(True), user_ids: str = Form(""),
                  db: Session = Depends(get_db), user: User = Depends(get_accounting_user)):
    s = get_settings(db)
    new_start = _d(start_date)
    has_journals = db.query(AccJournal.id).filter(AccJournal.source_key.isnot(None)).first() is not None
    if new_start != s.start_date and has_journals:
        raise HTTPException(400, "Clear the ledger before changing the start date")
    s.start_date = new_start
    s.auto_sync = auto_sync
    if user_ids:
        ids = [int(x) for x in json.loads(user_ids)]
        if user.id not in ids:
            raise HTTPException(400, "You cannot remove your own access")
        s.user_ids = json.dumps(ids)
    db.commit()
    return {"status": "saved"}


@router.post("/sync")
def run_sync(db: Session = Depends(get_db), user: User = Depends(get_accounting_user)):
    return sync_ledger(db, user.id)


@router.post("/reset")
def reset_ledger(confirm: str = Form(...), keep_manual: bool = Form(False),
                 db: Session = Depends(get_db), user: User = Depends(get_accounting_user)):
    """Test period only: removes journals so the ledger can restart from a new start date."""
    if confirm != "CLEAR":
        raise HTTPException(400, "Type CLEAR to confirm")
    q = db.query(AccJournal)
    if keep_manual:
        q = q.filter(AccJournal.source.notin_(("manual", "opening_manual")))
    ids = [j.id for j in q.all()]
    if ids:
        db.query(AccJournal).filter(AccJournal.id.in_(ids)).update({AccJournal.reversal_of: None, AccJournal.reversed_by: None},
                                                                  synchronize_session=False)
        db.query(AccJournalLine).filter(AccJournalLine.journal_id.in_(ids)).delete(synchronize_session=False)
        db.query(AccJournal).filter(AccJournal.id.in_(ids)).delete(synchronize_session=False)
    db.commit()
    return {"deleted": len(ids)}


# ---------- chart of accounts ----------

@router.get("/accounts")
def list_accounts(brand_id: int, db: Session = Depends(get_db), user: User = Depends(get_accounting_user)):
    ensure_chart(db)
    bal = _sums(db, brand_id)
    out = []
    for a in _accounts(db, brand_id):
        d, c = bal.get(a.id, (0.0, 0.0))
        j = _acc_json(a)
        j["balance"] = _r(d - c if a.type in DEBIT_NATURE else c - d)
        out.append(j)
    return out


@router.post("/accounts")
def create_account(brand_id: int = Form(...), code: str = Form(...), name: str = Form(...), name_ar: str = Form(""),
                   parent_code: str = Form(...), concept: str = Form(""),
                   db: Session = Depends(get_db), user: User = Depends(get_accounting_user)):
    _brand(db, brand_id)
    parent = db.query(AccAccount).filter(AccAccount.brand_id == brand_id, AccAccount.code == parent_code).first()
    if not parent:
        raise HTTPException(400, "Parent account not found")
    if db.query(AccAccount).filter(AccAccount.brand_id == brand_id, AccAccount.code == code.strip()).first():
        raise HTTPException(400, "Account code already exists")
    a = AccAccount(brand_id=brand_id, code=code.strip(), name=name.strip(), name_ar=name_ar.strip() or None, type=parent.type,
                   subtype=parent.subtype, parent_code=parent_code, concept=concept or parent.concept)
    db.add(a)
    db.commit()
    return _acc_json(a)


@router.put("/accounts/{account_id}")
def update_account(account_id: int, name: str = Form(...), name_ar: str = Form(""), concept: str = Form(""),
                   is_active: bool = Form(True), db: Session = Depends(get_db), user: User = Depends(get_accounting_user)):
    a = db.query(AccAccount).filter(AccAccount.id == account_id).first()
    if not a:
        raise HTTPException(404, "Account not found")
    a.name, a.name_ar, a.concept, a.is_active = name.strip(), name_ar.strip() or None, concept or a.concept, is_active
    db.commit()
    return _acc_json(a)


# ---------- account mapping ----------

@router.get("/mappings")
def list_mappings(brand_id: int, db: Session = Depends(get_db), user: User = Depends(get_accounting_user)):
    ensure_chart(db)
    L = Ledger(db)
    codes = {a.id: a for a in _accounts(db, brand_id)}

    def acc(aid):
        a = codes.get(aid)
        return {"account_id": aid, "code": a.code if a else None, "name": a.name if a else None, "name_ar": a.name_ar if a else None}

    rows = []
    for key, (code, label, label_ar) in SETTING_DEFAULTS.items():
        rows.append({"kind": "setting", "ref_id": key, "label": label, "label_ar": label_ar,
                     "mapped": (brand_id, "setting", key) in L.maps, **acc(L.setting(brand_id, key))})
    for c in db.query(ExpenseCategory).order_by(ExpenseCategory.name).all():
        rows.append({"kind": "expense_category", "ref_id": str(c.id), "label": c.name, "label_ar": c.name_ar or c.name,
                     "mapped": (brand_id, "expense_category", str(c.id)) in L.maps,
                     **acc(L.mapped(brand_id, "expense_category", c.id, c.name, EXPENSE_CATEGORY_DEFAULTS, "6270"))})
    for c in db.query(PurchaseCategory).order_by(PurchaseCategory.name).all():
        rows.append({"kind": "purchase_category", "ref_id": str(c.id), "label": c.name, "label_ar": c.name_ar or c.name,
                     "mapped": (brand_id, "purchase_category", str(c.id)) in L.maps,
                     **acc(L.mapped(brand_id, "purchase_category", c.id, c.name, PURCHASE_CATEGORY_DEFAULTS, "5110"))})
    db.commit()
    return rows


@router.post("/mappings")
def save_mapping(brand_id: int = Form(...), kind: str = Form(...), ref_id: str = Form(...), account_id: int = Form(0),
                 db: Session = Depends(get_db), user: User = Depends(get_accounting_user)):
    if kind not in ("setting", "expense_category", "purchase_category"):
        raise HTTPException(400, "Invalid mapping")
    m = db.query(AccMapping).filter(AccMapping.brand_id == brand_id, AccMapping.kind == kind, AccMapping.ref_id == ref_id).first()
    if not account_id:
        if m:
            db.delete(m)
        db.commit()
        return {"status": "default"}
    a = db.query(AccAccount).filter(AccAccount.id == account_id, AccAccount.brand_id == brand_id).first()
    if not a or a.is_header:
        raise HTTPException(400, "Choose a posting account of this company")
    if not m:
        m = AccMapping(brand_id=brand_id, kind=kind, ref_id=ref_id, account_id=account_id)
        db.add(m)
    m.account_id = account_id
    db.commit()
    return {"status": "saved"}


# ---------- journals ----------

def _journal_json(db: Session, j: AccJournal, accs: dict, with_lines: bool = True) -> dict:
    out = {"id": j.id, "number": j.number, "date": str(j.date), "description": j.description, "source": j.source,
           "source_key": j.source_key, "is_active": bool(j.is_active), "reversal_of": j.reversal_of, "reversed_by": j.reversed_by,
           "notes": j.notes}
    lines = db.query(AccJournalLine).filter(AccJournalLine.journal_id == j.id).order_by(AccJournalLine.id).all()
    out["total"] = _r(sum(l.debit or 0 for l in lines))
    if with_lines:
        out["lines"] = [{"account_id": l.account_id, "code": accs[l.account_id].code if l.account_id in accs else "",
                         "name": accs[l.account_id].name if l.account_id in accs else "",
                         "name_ar": accs[l.account_id].name_ar if l.account_id in accs else "",
                         "branch_id": l.branch_id, "debit": l.debit, "credit": l.credit, "memo": l.memo, "party": l.party} for l in lines]
    return out


@router.get("/journals")
def list_journals(brand_id: int, date_from: Optional[str] = None, date_to: Optional[str] = None, source: Optional[str] = None,
                  q: Optional[str] = None, page: int = 1, page_size: int = 100,
                  db: Session = Depends(get_db), user: User = Depends(get_accounting_user)):
    qry = db.query(AccJournal).filter(AccJournal.brand_id == brand_id)
    if date_from:
        qry = qry.filter(AccJournal.date >= _d(date_from))
    if date_to:
        qry = qry.filter(AccJournal.date <= _d(date_to))
    if source:
        qry = qry.filter(AccJournal.source == source)
    if q:
        qry = qry.filter(AccJournal.description.ilike(f"%{q}%") | AccJournal.source_key.ilike(f"%{q}%"))
    total = qry.count()
    rows = qry.order_by(AccJournal.date.desc(), AccJournal.number.desc()).offset((max(page, 1) - 1) * page_size).limit(page_size).all()
    accs = {a.id: a for a in _accounts(db, brand_id)}
    return {"total": total, "rows": [_journal_json(db, j, accs) for j in rows]}


@router.post("/journals")
def create_journal(brand_id: int = Form(...), date_: str = Form(..., alias="date"), description: str = Form(...),
                   lines: str = Form(...), notes: str = Form(""), opening: bool = Form(False),
                   db: Session = Depends(get_db), user: User = Depends(get_accounting_user)):
    _brand(db, brand_id)
    accs = {a.id: a for a in _accounts(db, brand_id)}
    parsed = []
    for ln in json.loads(lines):
        aid = int(ln["account_id"])
        if aid not in accs or accs[aid].is_header:
            raise HTTPException(400, "Each line needs a posting account of this company")
        dr, cr = _r(ln.get("debit")), _r(ln.get("credit"))
        if dr < 0 or cr < 0 or (dr > EPS and cr > EPS):
            raise HTTPException(400, "A line has either a debit or a credit")
        if dr > EPS or cr > EPS:
            parsed.append((aid, int(ln["branch_id"]) if ln.get("branch_id") else None, dr, cr, ln.get("memo") or "",
                           _valid_party(db, ln.get("party"))))
    if len(parsed) < 2:
        raise HTTPException(400, "A journal needs at least two lines")
    if abs(sum(p[2] - p[3] for p in parsed)) > EPS:
        raise HTTPException(400, "Debits and credits must be equal")
    j = post_journal(db, brand_id, _d(date_), description.strip(), parsed, "opening_manual" if opening else "manual",
                     user_id=user.id, notes=notes or None)
    db.commit()
    return _journal_json(db, j, accs)


@router.post("/journals/{journal_id}/reverse")
def reverse(journal_id: int, reason: str = Form(""), reverse_date: str = Form(""),
            db: Session = Depends(get_db), user: User = Depends(get_accounting_user)):
    j = db.query(AccJournal).filter(AccJournal.id == journal_id).first()
    if not j:
        raise HTTPException(404, "Journal not found")
    if j.source_key:
        raise HTTPException(400, "Automatic journals follow their Mudawwarah record; change or delete the record instead")
    if not j.is_active or j.reversal_of:
        raise HTTPException(400, "This journal is already reversed")
    rev = reverse_journal(db, j, reason or "Reversed", user.id, _d(reverse_date) or j.date)
    db.commit()
    return {"id": rev.id, "number": rev.number}


# ---------- reports ----------

def trial_balance_data(db: Session, brand_id: int, date_from: Optional[date], date_to: Optional[date], branch_id: Optional[int]):
    accs = _accounts(db, brand_id)
    opening = _sums(db, brand_id, date_from, branch_id=branch_id, before=True) if date_from else {}
    period = _sums(db, brand_id, date_from, date_to, branch_id)
    rows, tot = [], {"opening_dr": 0.0, "opening_cr": 0.0, "debit": 0.0, "credit": 0.0, "closing_dr": 0.0, "closing_cr": 0.0}
    for a in accs:
        od, oc = opening.get(a.id, (0.0, 0.0))
        pd_, pc = period.get(a.id, (0.0, 0.0))
        if max(abs(od), abs(oc), abs(pd_), abs(pc)) < EPS:
            continue
        ob = od - oc
        cb = ob + pd_ - pc
        r = {"account_id": a.id, "code": a.code, "name": a.name, "name_ar": a.name_ar, "type": a.type,
             "opening_dr": _r(max(ob, 0)), "opening_cr": _r(max(-ob, 0)), "debit": _r(pd_), "credit": _r(pc),
             "closing_dr": _r(max(cb, 0)), "closing_cr": _r(max(-cb, 0))}
        for k in tot:
            tot[k] += r[k]
        rows.append(r)
    return {"rows": rows, "totals": {k: _r(v) for k, v in tot.items()},
            "balanced": abs(tot["closing_dr"] - tot["closing_cr"]) < 0.01}


@router.get("/trial-balance")
def trial_balance(brand_id: int, date_from: Optional[str] = None, date_to: Optional[str] = None, branch_id: Optional[int] = None,
                  db: Session = Depends(get_db), user: User = Depends(get_accounting_user)):
    return trial_balance_data(db, brand_id, _d(date_from), _d(date_to), branch_id)


def ledger_data(db: Session, account_id: int, date_from: Optional[date], date_to: Optional[date], branch_id: Optional[int]):
    a = db.query(AccAccount).filter(AccAccount.id == account_id).first()
    if not a:
        raise HTTPException(404, "Account not found")
    sign = 1 if a.type in DEBIT_NATURE else -1
    base = db.query(AccJournalLine, AccJournal).join(AccJournal, AccJournal.id == AccJournalLine.journal_id)\
        .filter(AccJournalLine.account_id == account_id)
    if branch_id:
        base = base.filter(AccJournalLine.branch_id == branch_id)
    opening = 0.0
    if date_from:
        for l, _ in base.filter(AccJournal.date < date_from).all():
            opening += sign * ((l.debit or 0) - (l.credit or 0))
        base = base.filter(AccJournal.date >= date_from)
    if date_to:
        base = base.filter(AccJournal.date <= date_to)
    bal = opening
    rows = []
    for l, j in base.order_by(AccJournal.date, AccJournal.number, AccJournalLine.id).all():
        bal += sign * ((l.debit or 0) - (l.credit or 0))
        rows.append({"journal_id": j.id, "number": j.number, "date": str(j.date), "description": j.description, "source": j.source,
                     "branch_id": l.branch_id, "memo": l.memo, "debit": l.debit, "credit": l.credit, "balance": _r(bal)})
    return {"account": _acc_json(a), "opening": _r(opening), "closing": _r(bal), "rows": rows,
            "debit": _r(sum(r["debit"] or 0 for r in rows)), "credit": _r(sum(r["credit"] or 0 for r in rows))}


@router.get("/ledger")
def account_ledger(account_id: int, date_from: Optional[str] = None, date_to: Optional[str] = None, branch_id: Optional[int] = None,
                   db: Session = Depends(get_db), user: User = Depends(get_accounting_user)):
    return ledger_data(db, account_id, _d(date_from), _d(date_to), branch_id)


# ---------- payables / receivables (party sub-ledgers) ----------

CUSTOMERS = {
    "knet": ("KNET", "كي نت"), "link": ("Link / online payments", "لينك والدفع الإلكتروني"), "talabat": ("Talabat", "طلبات"),
    "keeta": ("Keeta", "كيتا"), "jahez": ("Jahez", "جاهز"), "snoonu": ("Snoonu", "سنونو"),
    "other": ("Other payment method (corporate customers)", "طريقة دفع أخرى (عملاء الشركات)"),
}
PARTY_KINDS = {"supplier": ("Supplier", "مورد"), "customer": ("Customer", "عميل"), "employee": ("Employee", "موظف"),
               "company": ("Related company", "شركة ذات علاقة")}


def _party_names(db: Session) -> dict[str, tuple[str, str]]:
    names = {f"customer:{k}": v for k, v in CUSTOMERS.items()}
    names.update({f"supplier:{x.id}": (x.name, x.name) for x in db.query(Supplier.id, Supplier.name).all()})
    names.update({f"employee:{x.id}": (f"{x.staff_no or x.id} - {x.name}", f"{x.staff_no or x.id} - {x.name_ar or x.name}")
                  for x in db.query(Employee).all()})
    names.update({f"company:{b.id}": (b.name_en, b.name_ar or b.name_en) for b in db.query(Brand).all()})
    return names


def _valid_party(db: Session, party: Optional[str]) -> Optional[str]:
    if not party:
        return None
    if party not in _party_names(db):
        raise HTTPException(400, f"Unknown party {party}")
    return party


def _party_json(names: dict, p: str) -> dict:
    kind = p.split(":", 1)[0]
    n = names.get(p, (p, p))
    k = PARTY_KINDS.get(kind, (kind, kind))
    return {"party": p, "kind": kind, "kind_name": k[0], "kind_name_ar": k[1], "name": n[0], "name_ar": n[1]}


@router.get("/party-options")
def party_options(db: Session = Depends(get_db), user: User = Depends(get_accounting_user)):
    names = _party_names(db)
    return sorted((_party_json(names, p) for p in names), key=lambda x: (x["kind"], x["name"].lower()))


def parties_data(db: Session, brand_id: int, side: str, as_of: Optional[date]) -> dict:
    """side 'payable' = liability accounts (credit balance), 'receivable' = asset accounts (debit balance)."""
    acc_type = "LIABILITY" if side == "payable" else "ASSET"
    sign = -1 if side == "payable" else 1
    accs = {a.id: a for a in _accounts(db, brand_id) if a.type == acc_type}
    q = db.query(AccJournalLine.party, AccJournalLine.account_id, func.sum(AccJournalLine.debit), func.sum(AccJournalLine.credit),
                 func.max(AccJournal.date)).join(AccJournal, AccJournal.id == AccJournalLine.journal_id)\
        .filter(AccJournal.brand_id == brand_id, AccJournalLine.account_id.in_(list(accs) or [0]))
    if as_of:
        q = q.filter(AccJournal.date <= as_of)
    names = _party_names(db)
    parties: dict[str, dict] = {}
    control: dict[int, dict] = {}
    for party, aid, d, c, last in q.group_by(AccJournalLine.party, AccJournalLine.account_id).all():
        amt = sign * ((d or 0) - (c or 0))
        a = accs[aid]
        ctl = control.setdefault(aid, {"account_id": aid, "code": a.code, "name": a.name, "name_ar": a.name_ar, "total": 0.0,
                                       "parties": 0.0, "used": False})
        ctl["total"] += amt
        if not party:
            continue
        ctl["used"] = True
        ctl["parties"] += amt
        r = parties.setdefault(party, {**_party_json(names, party), "debit": 0.0, "credit": 0.0, "balance": 0.0, "last_date": None,
                                       "accounts": []})
        r["debit"] = _r(r["debit"] + (d or 0))
        r["credit"] = _r(r["credit"] + (c or 0))
        r["balance"] = _r(r["balance"] + amt)
        r["last_date"] = max(filter(None, [r["last_date"], str(last)]))
        if a.code not in r["accounts"]:
            r["accounts"].append(a.code)
    rows = sorted(parties.values(), key=lambda x: (x["kind"], x["name"].lower()))
    ctl_rows = []
    for ctl in sorted(control.values(), key=lambda x: x["code"]):
        if not ctl["used"]:
            continue
        ctl["total"], ctl["parties"] = _r(ctl["total"]), _r(ctl["parties"])
        ctl["unallocated"] = _r(ctl["total"] - ctl["parties"])
        del ctl["used"]
        ctl_rows.append(ctl)
    return {"side": side, "as_of": str(as_of) if as_of else None, "rows": rows, "total": _r(sum(r["balance"] for r in rows)),
            "control": ctl_rows, "matched": all(abs(c["unallocated"]) < EPS for c in ctl_rows)}


@router.get("/parties")
def parties(brand_id: int, side: str = "payable", as_of: Optional[str] = None,
            db: Session = Depends(get_db), user: User = Depends(get_accounting_user)):
    if side not in ("payable", "receivable"):
        raise HTTPException(400, "side is payable or receivable")
    _brand(db, brand_id)
    return parties_data(db, brand_id, side, _d(as_of))


def party_ledger_data(db: Session, brand_id: int, party: str, date_from: Optional[date], date_to: Optional[date]) -> dict:
    accs = {a.id: a for a in _accounts(db, brand_id)}
    base = db.query(AccJournalLine, AccJournal).join(AccJournal, AccJournal.id == AccJournalLine.journal_id)\
        .filter(AccJournal.brand_id == brand_id, AccJournalLine.party == party)
    lines = base.order_by(AccJournal.date, AccJournal.number, AccJournalLine.id).all()
    payable = any(accs[l.account_id].type == "LIABILITY" for l, _ in lines if l.account_id in accs)
    sign = -1 if payable else 1
    opening = 0.0
    bal = 0.0
    rows = []
    for l, j in lines:
        amt = sign * ((l.debit or 0) - (l.credit or 0))
        if date_from and j.date < date_from:
            opening += amt
            bal = opening
            continue
        if date_to and j.date > date_to:
            continue
        bal += amt
        a = accs.get(l.account_id)
        rows.append({"journal_id": j.id, "number": j.number, "date": str(j.date), "description": j.description, "source": j.source,
                     "is_active": j.is_active, "code": a.code if a else "", "branch_id": l.branch_id, "memo": l.memo,
                     "debit": l.debit, "credit": l.credit, "balance": _r(bal)})
    return {**_party_json(_party_names(db), party), "side": "payable" if payable else "receivable", "opening": _r(opening),
            "closing": _r(bal), "rows": rows, "debit": _r(sum(r["debit"] or 0 for r in rows)),
            "credit": _r(sum(r["credit"] or 0 for r in rows))}


@router.get("/party-ledger")
def party_ledger(brand_id: int, party: str, date_from: Optional[str] = None, date_to: Optional[str] = None,
                 db: Session = Depends(get_db), user: User = Depends(get_accounting_user)):
    _brand(db, brand_id)
    return party_ledger_data(db, brand_id, party, _d(date_from), _d(date_to))


def _sections(accs: list[AccAccount], sums: dict, types: tuple) -> list[dict]:
    headers = {a.code: a for a in accs if a.is_header}
    groups: dict[str, dict] = {}
    for a in accs:
        if a.is_header or a.type not in types:
            continue
        d, c = sums.get(a.id, (0.0, 0.0))
        amt = (d - c) if a.type in DEBIT_NATURE else (c - d)
        if abs(amt) < EPS:
            continue
        parent = a.parent_code or a.code
        while parent in headers and headers[parent].parent_code and headers[parent].parent_code in headers \
                and headers[headers[parent].parent_code].parent_code:
            parent = headers[parent].parent_code
        h = headers.get(parent)
        g = groups.setdefault(parent, {"code": parent, "name": h.name if h else parent, "name_ar": h.name_ar if h else parent,
                                       "type": a.type, "accounts": [], "total": 0.0})
        g["accounts"].append({"account_id": a.id, "code": a.code, "name": a.name, "name_ar": a.name_ar, "amount": _r(amt)})
        g["total"] = _r(g["total"] + amt)
    return [groups[k] for k in sorted(groups)]


def income_statement_data(db: Session, brand_id: int, date_from: Optional[date], date_to: Optional[date], branch_id: Optional[int]):
    accs = _accounts(db, brand_id)
    sums = _sums(db, brand_id, date_from, date_to, branch_id)
    rev = _sections(accs, sums, ("REVENUE",))
    exp = _sections(accs, sums, ("EXPENSE",))
    revenue = _r(sum(g["total"] for g in rev))
    cos = _r(sum(g["total"] for g in exp if g["code"] == "5000"))
    opex = _r(sum(g["total"] for g in exp if g["code"] != "5000"))
    return {"revenue": rev, "expenses": exp, "total_revenue": revenue, "cost_of_sales": cos, "gross_profit": _r(revenue - cos),
            "operating_expenses": opex, "net_profit": _r(revenue - cos - opex)}


@router.get("/income-statement")
def income_statement(brand_id: int, date_from: Optional[str] = None, date_to: Optional[str] = None, branch_id: Optional[int] = None,
                     db: Session = Depends(get_db), user: User = Depends(get_accounting_user)):
    return income_statement_data(db, brand_id, _d(date_from), _d(date_to), branch_id)


def balance_sheet_data(db: Session, brand_id: int, as_of: date):
    accs = _accounts(db, brand_id)
    sums = _sums(db, brand_id, None, as_of)
    assets = _sections(accs, sums, ("ASSET",))
    liab = _sections(accs, sums, ("LIABILITY",))
    eq = _sections(accs, sums, ("EQUITY",))
    pl = income_statement_data(db, brand_id, None, as_of, None)["net_profit"]
    ta = _r(sum(g["total"] for g in assets))
    tl = _r(sum(g["total"] for g in liab))
    te = _r(sum(g["total"] for g in eq) + pl)
    return {"as_of": str(as_of), "assets": assets, "liabilities": liab, "equity": eq, "profit_to_date": pl, "total_assets": ta,
            "total_liabilities": tl, "total_equity": te, "balanced": abs(ta - tl - te) < 0.01}


@router.get("/balance-sheet")
def balance_sheet(brand_id: int, as_of: Optional[str] = None, db: Session = Depends(get_db), user: User = Depends(get_accounting_user)):
    return balance_sheet_data(db, brand_id, _d(as_of) or date.today())


def reconciliation_data(db: Session, brand_id: int, date_from: date, date_to: date):
    """Ledger figures next to Mudawwarah's own screens, so any difference shows the same day."""
    from app.routes.cash import _box_summary
    from app.routes.channels import channel_balance

    rows = []
    accs = {a.system_key: a for a in _accounts(db, brand_id) if a.system_key}
    upto = _sums(db, brand_id, None, date_to)
    for b in db.query(Branch).filter(Branch.brand_id == brand_id).order_by(Branch.id).all():
        a = accs.get(f"cashbox:{b.id}")
        d, c = upto.get(a.id, (0.0, 0.0)) if a else (0.0, 0.0)
        system = _box_summary(db, b.id, date_to)["closing_balance"]
        rows.append({"area": "cash_box", "name": b.name, "name_ar": b.name_ar, "system": _r(system), "ledger": _r(d - c),
                     "difference": _r(system - (d - c))})
    for ch in db.query(PaymentChannel).filter(PaymentChannel.brand_id == brand_id).all():
        a = accs.get(f"channel:{ch.id}")
        d, c = upto.get(a.id, (0.0, 0.0)) if a else (0.0, 0.0)
        system = channel_balance(db, ch, date_to)
        rows.append({"area": "channel", "name": ch.name, "name_ar": ch.name_ar, "system": _r(system), "ledger": _r(d - c),
                     "difference": _r(system - (d - c))})
    branch_ids = [b.id for b in db.query(Branch.id).filter(Branch.brand_id == brand_id).all()]
    cols = [Sale.physical_cash, Sale.physical_knet, Sale.physical_link, Sale.physical_wamd, Sale.physical_talabat, Sale.physical_keeta,
            Sale.physical_jahez, Sale.physical_other, Sale.physical_snoonu]
    sales = sum(float(db.query(func.coalesce(func.sum(col), 0)).filter(Sale.branch_id.in_(branch_ids or [0]), Sale.date >= date_from,
                                                                       Sale.date <= date_to).scalar()) for col in cols)
    period = _sums(db, brand_id, date_from, date_to)
    sale_led = 0.0
    for j_sum in db.query(func.coalesce(func.sum(AccJournalLine.credit - AccJournalLine.debit), 0))\
            .join(AccJournal, AccJournal.id == AccJournalLine.journal_id)\
            .join(AccAccount, AccAccount.id == AccJournalLine.account_id)\
            .filter(AccJournal.brand_id == brand_id, AccJournal.source == "sale", AccAccount.type == "REVENUE",
                    AccJournal.date >= date_from, AccJournal.date <= date_to).all():
        sale_led = float(j_sum[0])
    rows.append({"area": "sales", "name": "Sales (Sales screen)", "name_ar": "المبيعات", "system": _r(sales), "ledger": _r(sale_led),
                 "difference": _r(sales - sale_led)})
    exp = float(db.query(func.coalesce(func.sum(Expense.amount), 0)).filter(Expense.branch_id.in_(branch_ids or [0]),
                                                                            Expense.date >= date_from, Expense.date <= date_to).scalar())
    exp_led = 0.0
    for r in db.query(func.coalesce(func.sum(AccJournalLine.debit - AccJournalLine.credit), 0))\
            .join(AccJournal, AccJournal.id == AccJournalLine.journal_id)\
            .join(AccAccount, AccAccount.id == AccJournalLine.account_id)\
            .filter(AccJournal.brand_id == brand_id, AccJournal.source.in_(("expense", "salary")), AccAccount.type == "EXPENSE",
                    AccJournal.date >= date_from, AccJournal.date <= date_to).all():
        exp_led = float(r[0])
    loan_part = 0.0
    for r in db.query(func.coalesce(func.sum(AccJournalLine.credit - AccJournalLine.debit), 0))\
            .join(AccJournal, AccJournal.id == AccJournalLine.journal_id)\
            .join(AccAccount, AccAccount.id == AccJournalLine.account_id)\
            .filter(AccJournal.brand_id == brand_id, AccJournal.source == "salary", AccAccount.code == "1270",
                    AccJournal.date >= date_from, AccJournal.date <= date_to).all():
        loan_part = float(r[0])
    rows.append({"area": "expenses", "name": "Expenses (Expenses screen, salaries net of advances)", "name_ar": "المصروفات",
                 "system": _r(exp), "ledger": _r(exp_led - loan_part), "difference": _r(exp - exp_led + loan_part)})
    tb = trial_balance_data(db, brand_id, None, date_to, None)
    del period
    return {"date_from": str(date_from), "date_to": str(date_to), "rows": rows, "trial_balance_balanced": tb["balanced"],
            "all_matched": all(abs(r["difference"]) < 0.01 for r in rows) and tb["balanced"]}


@router.get("/reconciliation")
def reconciliation(brand_id: int, date_from: Optional[str] = None, date_to: Optional[str] = None,
                   db: Session = Depends(get_db), user: User = Depends(get_accounting_user)):
    s = get_settings(db)
    df = _d(date_from) or s.start_date or date.today()
    return reconciliation_data(db, brand_id, df, _d(date_to) or date.today())


# ---------- Excel ----------

def _xlsx(title: str, subtitle: str, header: list[str], rows: list[list], rtl: bool) -> StreamingResponse:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = Workbook()
    ws = wb.active
    ws.title = title[:30]
    ws.sheet_view.rightToLeft = rtl
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([subtitle])
    ws.append([])
    ws.append(header)
    for c in ws[4]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
        c.alignment = Alignment(horizontal="center")
    for r in rows:
        ws.append(r)
        if r and isinstance(r[0], str) and r[0].startswith("**"):
            ws.cell(ws.max_row, 1).value = r[0][2:]
            for c in ws[ws.max_row]:
                c.font = Font(bold=True)
    for col in ws.columns:
        width = max((len(str(c.value)) for c in col if c.value is not None and c.row > 3), default=8)
        ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 10), 60)
    for row in ws.iter_rows(min_row=5):
        for c in row:
            if isinstance(c.value, float):
                c.number_format = "#,##0.000"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = title.replace(" ", "_") + ".xlsx"
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="{fname}"'})


@router.get("/export/{report}")
def export_report(report: str, brand_id: int, lang: str = "en", date_from: Optional[str] = None, date_to: Optional[str] = None,
                  as_of: Optional[str] = None, account_id: Optional[int] = None, branch_id: Optional[int] = None,
                  side: str = "payable", party: Optional[str] = None, db: Session = Depends(get_db), user: User = Depends(get_accounting_user)):
    ar = lang == "ar"
    b = _brand(db, brand_id)
    company = (b.name_ar or b.name_en) if ar else b.name_en
    nm = (lambda o: (o.get("name_ar") or o.get("name")) if ar else o.get("name"))
    period = f"{date_from or ''} - {date_to or ''}"
    if report == "trial-balance":
        d = trial_balance_data(db, brand_id, _d(date_from), _d(date_to), branch_id)
        hdr = ["الرمز", "الحساب", "افتتاحي مدين", "افتتاحي دائن", "مدين", "دائن", "ختامي مدين", "ختامي دائن"] if ar else \
            ["Code", "Account", "Opening Dr", "Opening Cr", "Debit", "Credit", "Closing Dr", "Closing Cr"]
        rows = [[r["code"], nm(r), r["opening_dr"], r["opening_cr"], r["debit"], r["credit"], r["closing_dr"], r["closing_cr"]]
                for r in d["rows"]]
        t = d["totals"]
        rows.append(["**" + ("الإجمالي" if ar else "Total"), "", t["opening_dr"], t["opening_cr"], t["debit"], t["credit"],
                     t["closing_dr"], t["closing_cr"]])
        return _xlsx(("ميزان المراجعة" if ar else "Trial Balance"), f"{company} {period}", hdr, rows, ar)
    if report == "ledger" and account_id:
        d = ledger_data(db, account_id, _d(date_from), _d(date_to), branch_id)
        hdr = ["التاريخ", "رقم القيد", "البيان", "مدين", "دائن", "الرصيد"] if ar else ["Date", "Journal", "Description", "Debit", "Credit", "Balance"]
        rows = [["", "", "رصيد افتتاحي" if ar else "Opening balance", None, None, d["opening"]]]
        rows += [[r["date"], r["number"], r["description"], r["debit"], r["credit"], r["balance"]] for r in d["rows"]]
        rows.append(["**" + ("الإجمالي" if ar else "Total"), "", "", d["debit"], d["credit"], d["closing"]])
        a = d["account"]
        return _xlsx(("كشف حساب" if ar else "Account Ledger"), f"{company} {a['code']} {nm(a)} {period}", hdr, rows, ar)
    if report == "parties" and side in ("payable", "receivable"):
        d = parties_data(db, brand_id, side, _d(as_of))
        pay = side == "payable"
        hdr = ["النوع", "الاسم", "الحسابات", "مدين", "دائن", "الرصيد", "آخر حركة"] if ar else \
            ["Type", "Name", "Accounts", "Debit", "Credit", "Balance", "Last entry"]
        rows = [[(r["kind_name_ar"] if ar else r["kind_name"]), nm(r), ", ".join(r["accounts"]), r["debit"], r["credit"], r["balance"],
                 r["last_date"]] for r in d["rows"]]
        rows.append(["**" + ("الإجمالي" if ar else "Total"), "", "", None, None, d["total"], ""])
        title = ("الذمم الدائنة" if pay else "الذمم المدينة") if ar else ("Payables" if pay else "Receivables")
        return _xlsx(title, f"{company} {as_of or ''}", hdr, rows, ar)
    if report == "party-ledger" and party:
        d = party_ledger_data(db, brand_id, party, _d(date_from), _d(date_to))
        hdr = ["التاريخ", "رقم القيد", "البيان", "الحساب", "مدين", "دائن", "الرصيد"] if ar else \
            ["Date", "Journal", "Description", "Account", "Debit", "Credit", "Balance"]
        rows = [["", "", "رصيد افتتاحي" if ar else "Opening balance", "", None, None, d["opening"]]]
        rows += [[r["date"], r["number"], r["description"], r["code"], r["debit"], r["credit"], r["balance"]] for r in d["rows"]]
        rows.append(["**" + ("الإجمالي" if ar else "Total"), "", "", "", d["debit"], d["credit"], d["closing"]])
        return _xlsx(("كشف حساب" if ar else "Statement of account"), f"{company} {nm(d)} {period}", hdr, rows, ar)
    if report == "income-statement":
        d = income_statement_data(db, brand_id, _d(date_from), _d(date_to), branch_id)
        rows = []
        for g in d["revenue"] + d["expenses"]:
            rows.append(["**" + nm(g), None])
            rows += [[f"   {a['code']} {nm(a)}", a["amount"]] for a in g["accounts"]]
            rows.append([("إجمالي " if ar else "Total ") + nm(g), g["total"]])
        rows += [["**" + ("مجمل الربح" if ar else "Gross profit"), d["gross_profit"]],
                 ["**" + ("صافي الربح" if ar else "Net profit"), d["net_profit"]]]
        return _xlsx(("قائمة الدخل" if ar else "Income Statement"), f"{company} {period}",
                     ["البند" if ar else "Item", "د.ك" if ar else "KWD"], rows, ar)
    if report == "balance-sheet":
        d = balance_sheet_data(db, brand_id, _d(as_of) or date.today())
        rows = []
        for title, groups, total in ((("الأصول" if ar else "Assets"), d["assets"], d["total_assets"]),
                                     (("الالتزامات" if ar else "Liabilities"), d["liabilities"], d["total_liabilities"]),
                                     (("حقوق الملكية" if ar else "Equity"), d["equity"], None)):
            rows.append(["**" + title, None])
            for g in groups:
                rows += [[f"   {a['code']} {nm(a)}", a["amount"]] for a in g["accounts"]]
            if total is not None:
                rows.append(["**" + ("إجمالي " if ar else "Total ") + title, total])
        rows.append(["   " + ("ربح الفترة" if ar else "Profit to date"), d["profit_to_date"]])
        rows.append(["**" + ("إجمالي حقوق الملكية" if ar else "Total equity"), d["total_equity"]])
        return _xlsx(("الميزانية العمومية" if ar else "Balance Sheet"), f"{company} {d['as_of']}",
                     ["البند" if ar else "Item", "د.ك" if ar else "KWD"], rows, ar)
    raise HTTPException(400, "Unknown report")
