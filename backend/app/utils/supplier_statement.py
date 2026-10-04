"""Central Purchases supplier Statement of Account (PDF / Excel) with the brand's name and logo."""
import io
import os
from datetime import date
from typing import Optional

import arabic_reshaper
from bidi.algorithm import get_display
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy.orm import Session

from app.database import UPLOAD_DIR
from app.models.channel import PaymentChannel
from app.models.hr import Brand
from app.models.procurement import ProcInvoice, ProcOrder, ProcPayment
from app.models.purchase import Supplier

LABELS = {
    "en": {
        "title": "STATEMENT OF ACCOUNT", "subtitle": "Central Purchases - Supplier Account",
        "supplier": "Supplier", "contact": "Contact", "terms": "Payment Terms", "period": "Period",
        "printed": "Printed On", "all_dates": "All dates", "to": "to", "cash": "Cash", "credit_terms": "Credit",
        "date": "Date", "type": "Type", "ref": "Invoice / Ref", "po": "PO No", "desc": "Description",
        "due": "Due Date", "debit": "Debit (Invoices)", "credit": "Credit (Payments)", "balance": "Balance",
        "bf": "Balance brought forward", "total": "Total", "summary": "Account Summary",
        "s_opening": "Opening Balance", "s_invoiced": "Total Invoiced", "s_paid": "Total Paid",
        "s_closing": "Closing Balance (Amount Due)", "aging": "Outstanding Ageing (as of {d})",
        "a_current": "Not yet due", "a_30": "1-30 days", "a_60": "31-60 days", "a_90": "61-90 days",
        "a_90p": "Over 90 days", "a_total": "Total Due",
        "k_opening": "Opening Balance", "k_invoice": "Purchase Invoice", "k_payment": "Payment",
        "against": "against", "no_entries": "No transactions in this period",
        "note": "Amounts in Kuwaiti Dinar (KD). Please review this statement and notify the Purchase Office of any discrepancy.",
        "prepared": "Prepared by", "approved": "Approved by", "confirm": "Supplier confirmation",
        "page": "Page",
    },
    "ar": {
        "title": "كشف حساب", "subtitle": "المشتريات المركزية - حساب المورد",
        "supplier": "المورد", "contact": "التواصل", "terms": "شروط الدفع", "period": "الفترة",
        "printed": "تاريخ الطباعة", "all_dates": "جميع التواريخ", "to": "إلى", "cash": "نقدي", "credit_terms": "آجل",
        "date": "التاريخ", "type": "النوع", "ref": "الفاتورة / المرجع", "po": "رقم أمر الشراء", "desc": "البيان",
        "due": "تاريخ الاستحقاق", "debit": "مدين (فواتير)", "credit": "دائن (مدفوعات)", "balance": "الرصيد",
        "bf": "رصيد مرحل", "total": "الإجمالي", "summary": "ملخص الحساب",
        "s_opening": "الرصيد الافتتاحي", "s_invoiced": "إجمالي الفواتير", "s_paid": "إجمالي المدفوع",
        "s_closing": "الرصيد الختامي (المبلغ المستحق)", "aging": "أعمار الديون المستحقة (حتى {d})",
        "a_current": "غير مستحق بعد", "a_30": "1-30 يوم", "a_60": "31-60 يوم", "a_90": "61-90 يوم",
        "a_90p": "أكثر من 90 يوم", "a_total": "إجمالي المستحق",
        "k_opening": "رصيد افتتاحي", "k_invoice": "فاتورة شراء", "k_payment": "دفعة",
        "against": "عن", "no_entries": "لا توجد حركات في هذه الفترة",
        "note": "المبالغ بالدينار الكويتي. يرجى مراجعة هذا الكشف وإبلاغ مكتب المشتريات بأي اختلاف.",
        "prepared": "أعده", "approved": "اعتمده", "confirm": "مصادقة المورد",
        "page": "صفحة",
    },
}

METHODS = {
    "en": {"purchase_petty_cash": "Purchase Office Petty Cash", "bank_transfer": "Bank Transfer",
           "knet": "KNET", "cheque": "Cheque"},
    "ar": {"purchase_petty_cash": "نثرية مكتب المشتريات", "bank_transfer": "تحويل بنكي",
           "knet": "كي نت", "cheque": "شيك"},
}

AGING_KEYS = ("a_current", "a_30", "a_60", "a_90", "a_90p")


def _d(s: Optional[str]) -> Optional[date]:
    return date.fromisoformat(s) if s else None


def logo_file(brand: Optional[Brand]) -> Optional[str]:
    if brand is None or not brand.logo_path:
        return None
    path = os.path.join(UPLOAD_DIR, brand.logo_path)
    return path if os.path.exists(path) else None


def build(db: Session, brand_id: int, supplier_id: int, date_from: Optional[str], date_to: Optional[str]) -> dict:
    brand = db.query(Brand).filter(Brand.id == brand_id).first()
    sup = db.query(Supplier).filter(Supplier.id == supplier_id).first()
    invs = db.query(ProcInvoice).filter(ProcInvoice.brand_id == brand_id, ProcInvoice.supplier_id == supplier_id).all()
    order_ids = [i.order_id for i in invs if i.order_id]
    orders = {o.id: o for o in db.query(ProcOrder).filter(ProcOrder.id.in_(order_ids)).all()} if order_ids else {}
    inv_by_id = {i.id: i for i in invs}
    pays = db.query(ProcPayment).filter(ProcPayment.invoice_id.in_(list(inv_by_id))).all() if inv_by_id else []
    channels = {c.id: c for c in db.query(PaymentChannel).all()}

    entries = []
    for i in invs:
        o = orders.get(i.order_id)
        opening = i.kind == "opening"
        entries.append({"date": i.date, "seq": 0 if opening else 1, "id": i.id,
                        "kind": "opening" if opening else "invoice",
                        "ref": "" if opening else (i.invoice_number or ""), "po_no": o.po_no if o else "",
                        "due": i.due_date, "notes": i.notes or "", "method": "", "channel": None, "against": "",
                        "debit": round(i.total_amount or 0, 3), "credit": 0.0})
    for p in pays:
        i = inv_by_id[p.invoice_id]
        o = orders.get(i.order_id)
        entries.append({"date": p.date, "seq": 2, "id": p.id, "kind": "payment", "ref": p.reference or "",
                        "po_no": o.po_no if o else "", "due": None, "notes": p.notes or "", "method": p.method or "",
                        "channel": channels.get(p.channel_id) if p.channel_id else None,
                        "against": "opening" if i.kind == "opening" else (i.invoice_number or (o.po_no if o else "")),
                        "debit": 0.0, "credit": round(p.amount or 0, 3)})
    entries.sort(key=lambda e: (e["date"], e["seq"], e["id"]))

    df, dt = _d(date_from), _d(date_to)
    opening_bal = round(sum(e["debit"] - e["credit"] for e in entries if df and e["date"] < df), 3)
    period = [e for e in entries if (not df or e["date"] >= df) and (not dt or e["date"] <= dt)]
    run = opening_bal
    for e in period:
        run += e["debit"] - e["credit"]
        e["balance"] = round(run, 3)

    as_of = dt or date.today()
    paid_by_inv: dict = {}
    for p in pays:
        if p.date <= as_of:
            paid_by_inv[p.invoice_id] = paid_by_inv.get(p.invoice_id, 0) + (p.amount or 0)
    aging = {k: 0.0 for k in AGING_KEYS}
    for i in invs:
        if i.date > as_of:
            continue
        due = round((i.total_amount or 0) - paid_by_inv.get(i.id, 0), 3)
        if due <= 0.0005:
            continue
        days = (as_of - (i.due_date or i.date)).days
        key = "a_current" if days <= 0 else "a_30" if days <= 30 else "a_60" if days <= 60 else "a_90" if days <= 90 else "a_90p"
        aging[key] += due

    return {
        "brand": brand, "supplier": sup, "date_from": df, "date_to": dt, "as_of": as_of,
        "opening": opening_bal, "entries": period,
        "invoiced": round(sum(e["debit"] for e in period), 3),
        "paid": round(sum(e["credit"] for e in period), 3),
        "closing": round(run, 3),
        "aging": {k: round(v, 3) for k, v in aging.items()},
    }


def _fmt_date(d: Optional[date]) -> str:
    return d.strftime("%d-%m-%Y") if d else ""


def _period(data: dict, L: dict) -> str:
    df, dt = data["date_from"], data["date_to"]
    if not df and not dt:
        return L["all_dates"]
    return f"{_fmt_date(df) or '...'} {L['to']} {_fmt_date(dt) or _fmt_date(date.today())}"


def _describe(e: dict, lang: str) -> str:
    L, M = LABELS[lang], METHODS[lang]
    parts = [L["k_" + e["kind"]]]
    if e["kind"] == "payment":
        parts.append(M.get(e["method"], e["method"]))
        ch = e["channel"]
        if ch is not None:
            parts.append((ch.name_ar or ch.name) if lang == "ar" else ch.name)
        if e["against"]:
            parts.append(f"{L['against']} {L['k_opening'] if e['against'] == 'opening' else e['against']}")
    if e["notes"]:
        parts.append(e["notes"])
    return " - ".join(p for p in parts if p)


def _rows(data: dict, lang: str) -> list:
    """Statement table rows (plain values) incl. brought-forward and total rows."""
    L = LABELS[lang]
    rows = [["", L["bf"] if data["date_from"] else L["s_opening"], "", "", "", "", None, None, data["opening"]]]
    for e in data["entries"]:
        rows.append([_fmt_date(e["date"]), L["k_" + e["kind"]], e["ref"], e["po_no"], _describe(e, lang),
                     _fmt_date(e["due"]), e["debit"] or None, e["credit"] or None, e["balance"]])
    rows.append(["", L["total"], "", "", "", "", data["invoiced"], data["paid"], data["closing"]])
    return rows


def _terms(sup: Optional[Supplier], L: dict) -> str:
    if sup is None:
        return ""
    return L["credit_terms"] if sup.payment_type == "credit" else L["cash"]


def _contact(sup: Optional[Supplier]) -> str:
    if sup is None:
        return ""
    return " / ".join(x for x in [sup.whatsapp, sup.email] if x) or "-"


def render_pdf(data: dict, lang: str) -> bytes:
    lang = "ar" if lang == "ar" else "en"
    L = LABELS[lang]
    rtl = lang == "ar"
    for name, file in (("DejaVuSans", "DejaVuSans.ttf"), ("DejaVuSans-Bold", "DejaVuSans-Bold.ttf")):
        if name not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont(name, f"/usr/share/fonts/truetype/dejavu/{file}"))

    def tx(s: str) -> str:
        s = str(s or "")
        return get_display(arabic_reshaper.reshape(s)) if s else ""

    start = 2 if rtl else 0
    end = 0 if rtl else 2
    st = lambda n, size=8, bold=False, align=start, color=colors.black: ParagraphStyle(
        n, fontName="DejaVuSans-Bold" if bold else "DejaVuSans", fontSize=size, leading=size + 3,
        alignment=align, textColor=color)
    cell, cellb = st("c", 7.5), st("cb", 7.5, True)
    num, numb = st("n", 7.5, align=end), st("nb", 7.5, True, align=end)
    hdr = st("h", size=8, bold=True, color=colors.white)
    hdrn = st("hn", size=8, bold=True, align=end, color=colors.white)
    navy = colors.HexColor("#1F3A5F")
    light = colors.HexColor("#EEF3F8")

    def money(v: Optional[float], bold: bool = False) -> Paragraph:
        return Paragraph("" if v is None else f"{v:,.3f}", numb if bold else num)

    def order(cols: list) -> list:
        return list(reversed(cols)) if rtl else cols

    brand, sup = data["brand"], data["supplier"]
    page_w = A4[0] - 24 * mm
    el = []

    title_block = [Paragraph(tx(brand.name_en if brand else ""), st("bn", 15, True, 1, navy)),
                   Paragraph(tx(brand.name_ar or "") if brand else "", st("ba", 13, True, 1, navy)),
                   Spacer(1, 2 * mm),
                   Paragraph(tx(L["title"]), st("t", 13, True, 1)),
                   Paragraph(tx(L["subtitle"]), st("s", 8, False, 1, colors.grey))]
    logo = logo_file(brand)
    if logo:
        iw, ih = ImageReader(logo).getSize()
        h = 24 * mm
        w = min(h * iw / ih, 45 * mm)
        img = Image(logo, width=w, height=w * ih / iw)
        head = Table([order([img, title_block, ""])], colWidths=order([48 * mm, page_w - 96 * mm, 48 * mm]))
    else:
        head = Table([[title_block]], colWidths=[page_w])
    head.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                              ("LINEBELOW", (0, 0), (-1, 0), 1.5, navy),
                              ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))
    el += [head, Spacer(1, 4 * mm)]

    info = [
        order([Paragraph(tx(L["supplier"]), cellb), Paragraph(tx(sup.name if sup else ""), st("sn", 10, True)),
               Paragraph(tx(L["period"]), cellb), Paragraph(tx(_period(data, L)), cell)]),
        order([Paragraph(tx(L["contact"]), cellb), Paragraph(tx(_contact(sup)), cell),
               Paragraph(tx(L["printed"]), cellb), Paragraph(_fmt_date(date.today()), cell)]),
        order([Paragraph(tx(L["terms"]), cellb), Paragraph(tx(_terms(sup, L)), cell),
               Paragraph(tx(L["s_closing"]), cellb), money(data["closing"], True)]),
    ]
    it = Table(info, colWidths=order([28 * mm, page_w / 2 - 28 * mm, 40 * mm, page_w / 2 - 40 * mm]))
    lab_cols = (3, 1) if rtl else (0, 2)
    it.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#B0BEC5")),
                            ("BACKGROUND", (lab_cols[0], 0), (lab_cols[0], -1), light),
                            ("BACKGROUND", (lab_cols[1], 0), (lab_cols[1], -1), light),
                            ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
    el += [it, Spacer(1, 5 * mm)]

    widths = [21 * mm, 17 * mm, 20 * mm, 24 * mm, 0, 21 * mm, 18 * mm, 18 * mm, 20 * mm]
    widths[4] = page_w - sum(widths)
    heads = [L[k] for k in ("date", "type", "ref", "po", "desc", "due", "debit", "credit", "balance")]
    table = [order([Paragraph(tx(h), hdrn if n >= 6 else hdr) for n, h in enumerate(heads)])]
    rows = _rows(data, lang)
    for r_i, r in enumerate(rows):
        edge = r_i in (0, len(rows) - 1)
        table.append(order([Paragraph(tx(v), cellb if edge else cell) if n < 6 else money(v, edge or n == 8)
                            for n, v in enumerate(r)]))
    if not data["entries"]:
        table.insert(2, order([Paragraph(tx(L["no_entries"]), cell)] + [""] * 8))
    t = Table(table, colWidths=order(widths), repeatRows=1)
    style = [("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#B0BEC5")),
             ("BACKGROUND", (0, 0), (-1, 0), navy),
             ("BACKGROUND", (0, 1), (-1, 1), light),
             ("BACKGROUND", (0, -1), (-1, -1), light),
             ("LINEABOVE", (0, -1), (-1, -1), 1, navy),
             ("LEFTPADDING", (0, 0), (-1, -1), 3), ("RIGHTPADDING", (0, 0), (-1, -1), 3),
             ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]
    for r_i in range(2, len(table) - 1, 2):
        style.append(("BACKGROUND", (0, r_i), (-1, r_i), colors.HexColor("#FAFBFC")))
    t.setStyle(TableStyle(style))
    el += [t, Spacer(1, 6 * mm)]

    summ = [[Paragraph(tx(L["summary"]), hdr), ""]]
    for k, v in (("s_opening", data["opening"]), ("s_invoiced", data["invoiced"]),
                 ("s_paid", data["paid"]), ("s_closing", data["closing"])):
        summ.append(order([Paragraph(tx(L[k]), cellb if k == "s_closing" else cell), money(v, k == "s_closing")]))
    sm = Table(summ, colWidths=order([50 * mm, 28 * mm]))
    sm.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#B0BEC5")),
                            ("SPAN", (0, 0), (1, 0)), ("BACKGROUND", (0, 0), (-1, 0), navy),
                            ("BACKGROUND", (0, -1), (-1, -1), light)]))
    ag_heads = [L[k] for k in AGING_KEYS] + [L["a_total"]]
    ag_vals = [data["aging"][k] for k in AGING_KEYS] + [round(sum(data["aging"].values()), 3)]
    ag = Table([[Paragraph(tx(L["aging"].format(d=_fmt_date(data["as_of"]))), hdr)] + [""] * 5,
                order([Paragraph(tx(h), st("ah", 7, True, 1)) for h in ag_heads]),
                order([money(v, n == 5) for n, v in enumerate(ag_vals)])],
               colWidths=[(page_w - 84 * mm) / 6] * 6)
    ag.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#B0BEC5")),
                            ("SPAN", (0, 0), (-1, 0)), ("BACKGROUND", (0, 0), (-1, 0), navy),
                            ("BACKGROUND", (0, 1), (-1, 1), light), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                            ("LEFTPADDING", (0, 0), (-1, -1), 3), ("RIGHTPADDING", (0, 0), (-1, -1), 3)]))
    box = Table([order([sm, "", ag])], colWidths=order([80 * mm, 4 * mm, page_w - 84 * mm]))
    box.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0),
                             ("RIGHTPADDING", (0, 0), (-1, -1), 0)]))
    el += [box, Spacer(1, 5 * mm), Paragraph(tx(L["note"]), st("nt", 7, False, start, colors.grey)), Spacer(1, 14 * mm)]
    sig = Table([order([Paragraph(tx(L[k]), st("sg", 8, False, 1)) for k in ("prepared", "approved", "confirm")]),
                 order([Paragraph("______________________", st("sl", 8, False, 1))] * 3)],
                colWidths=[page_w / 3] * 3)
    el.append(sig)

    footer_left = tx(f"{brand.name_en if brand else ''} - {L['title']} - {sup.name if sup else ''}")

    def on_page(canvas, doc):
        canvas.saveState()
        canvas.setFont("DejaVuSans", 7)
        canvas.setFillColor(colors.grey)
        canvas.drawString(12 * mm, 8 * mm, footer_left)
        canvas.drawRightString(A4[0] - 12 * mm, 8 * mm, f"{tx(L['page'])} {doc.page}")
        canvas.restoreState()

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=12 * mm, rightMargin=12 * mm,
                            topMargin=12 * mm, bottomMargin=15 * mm,
                            title=f"Statement of Account - {sup.name if sup else ''}")
    doc.build(el, onFirstPage=on_page, onLaterPages=on_page)
    return buf.getvalue()


def render_xlsx(data: dict, lang: str) -> bytes:
    lang = "ar" if lang == "ar" else "en"
    L = LABELS[lang]
    brand, sup = data["brand"], data["supplier"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Statement"
    ws.sheet_view.rightToLeft = lang == "ar"
    navy = PatternFill("solid", fgColor="1F3A5F")
    light = PatternFill("solid", fgColor="EEF3F8")
    thin = Side(style="thin", color="B0BEC5")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for col, w in zip("ABCDEFGHI", (12, 16, 16, 16, 48, 12, 16, 16, 16)):
        ws.column_dimensions[col].width = w

    logo = logo_file(brand)
    if logo:
        img = XLImage(logo)
        ratio = 70 / img.height if img.height else 1
        img.height, img.width = 70, int(img.width * ratio)
        ws.add_image(img, "A1")
    for r, (text, size) in enumerate(((brand.name_en if brand else "", 16), (brand.name_ar or "" if brand else "", 14),
                                      (L["title"], 13), (L["subtitle"], 9)), 1):
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=7)
        c = ws.cell(row=r, column=3, value=text)
        c.font = Font(bold=r < 4, size=size, color="1F3A5F" if r < 3 else "000000")
        c.alignment = Alignment(horizontal="center")

    info = [(L["supplier"], sup.name if sup else "", L["period"], _period(data, L)),
            (L["contact"], _contact(sup), L["printed"], _fmt_date(date.today())),
            (L["terms"], _terms(sup, L), L["s_closing"], data["closing"])]
    for i, (a, b, c, d) in enumerate(info, 6):
        for col, v, lab in ((1, a, True), (2, b, False), (6, c, True), (8, d, False)):
            cell = ws.cell(row=i, column=col, value=v)
            cell.font = Font(bold=lab or (col == 8 and i == 8))
            if lab:
                cell.fill = light
        ws.merge_cells(start_row=i, start_column=2, end_row=i, end_column=5)
        ws.merge_cells(start_row=i, start_column=6, end_row=i, end_column=7)
        if isinstance(d, float):
            ws.cell(row=i, column=8).number_format = "#,##0.000"

    hr = 10
    for n, k in enumerate(("date", "type", "ref", "po", "desc", "due", "debit", "credit", "balance"), 1):
        c = ws.cell(row=hr, column=n, value=L[k])
        c.font, c.fill, c.border = Font(bold=True, color="FFFFFF"), navy, border
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    rows = _rows(data, lang)
    for r_i, r in enumerate(rows):
        edge = r_i in (0, len(rows) - 1)
        for n, v in enumerate(r, 1):
            c = ws.cell(row=hr + 1 + r_i, column=n, value=v)
            c.border = border
            if n >= 7:
                c.number_format = "#,##0.000"
            if edge:
                c.font, c.fill = Font(bold=True), light
            if n == 5:
                c.alignment = Alignment(wrap_text=True, vertical="top")

    r = hr + len(rows) + 2
    ws.cell(row=r, column=1, value=L["summary"]).font = Font(bold=True, size=11)
    for k, v in (("s_opening", data["opening"]), ("s_invoiced", data["invoiced"]),
                 ("s_paid", data["paid"]), ("s_closing", data["closing"])):
        r += 1
        ws.cell(row=r, column=1, value=L[k]).font = Font(bold=k == "s_closing")
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        c = ws.cell(row=r, column=3, value=v)
        c.number_format, c.font = "#,##0.000", Font(bold=k == "s_closing")
    r += 2
    ws.cell(row=r, column=1, value=L["aging"].format(d=_fmt_date(data["as_of"]))).font = Font(bold=True, size=11)
    r += 1
    vals = [data["aging"][k] for k in AGING_KEYS] + [round(sum(data["aging"].values()), 3)]
    for n, (h, v) in enumerate(zip([L[k] for k in AGING_KEYS] + [L["a_total"]], vals), 1):
        hc = ws.cell(row=r, column=n, value=h)
        hc.font, hc.fill, hc.border = Font(bold=True), light, border
        vc = ws.cell(row=r + 1, column=n, value=v)
        vc.number_format, vc.border = "#,##0.000", border
    ws.cell(row=r + 3, column=1, value=L["note"]).font = Font(italic=True, size=8, color="808080")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
